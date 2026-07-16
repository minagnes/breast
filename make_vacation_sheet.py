# -*- coding: utf-8 -*-
"""
휴가/부재 입력용 엑셀(vacation.xlsx) 생성기
================================================
schedule_generator.py 가 쓰는 것과 같은 달력(월~토, 그 주의 수요일이 해당 월에
속하는 주만 포함)을 기준으로, 왼쪽에 근무자 목록 / 위쪽에 날짜를 배치한
빈 입력표를 만든다. 여기에 휴가·반차·학회·출장 등을 적어 넣은 뒤,
schedule_generator.py 의 VACATIONS 리스트를 채울 때 참고하면 된다.

사용법:
    python make_vacation_sheet.py
    → 같은 폴더에 vacation.xlsx 생성 (이미 있으면 vacation-1.xlsx 로 저장)
"""

import os
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

# ============================================================
# 설정 - schedule_generator.py 와 동일하게 맞출 것
# ============================================================

YEAR = 2026
MONTH = 9

# 추석 연휴 (2026년 9/24(목)~9/26(토))
HOLIDAYS = {
    date(2026, 9, 24): "추석",
    date(2026, 9, 25): "추석",
    date(2026, 9, 26): "추석",
}

# 근무자 목록 (2026년 9월 기준: K3 퇴사 → K5 대체, F4 신규 / 스케줄_배정_규정.docx 2장)
K_STAFF = ["K1", "K2", "K4", "K5"]
F_STAFF = ["F1", "F2", "F3", "F4"]
J_STAFF = ["J1", "J2"]
R_STAFF = ["R"]

STAFF_LIST = K_STAFF + F_STAFF + J_STAFF + R_STAFF

WEEKDAY_NAMES = ["월", "화", "수", "목", "금", "토"]

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))

# 입력 예시로 제공할 드롭다운 선택지 (자유 텍스트 입력도 가능)
CHOICES = ["", "휴가(종일)", "휴가(오전)", "휴가(오후)", "학회(종일)", "학회(오전)", "학회(오후)",
           "출장(종일)", "출장(오전)", "출장(오후)", "교육(종일)", "반차(오전)", "반차(오후)"]

# ============================================================
# 달력 유틸 (schedule_generator.py 와 동일한 규칙)
# ============================================================


def month_weeks(year, month):
    """그 주의 수요일이 해당 월에 속하는 모든 주의 월요일 리스트."""
    probe = date(year, month, 15) - timedelta(days=40)
    mondays = []
    d = probe - timedelta(days=probe.weekday())
    for _ in range(16):
        wed = d + timedelta(days=2)
        if wed.year == year and wed.month == month:
            mondays.append(d)
        d += timedelta(days=7)
    return sorted(set(mondays))


def all_dates(year, month):
    dates = []
    for monday in month_weeks(year, month):
        dates.extend(monday + timedelta(days=i) for i in range(6))  # 월~토
    return dates


# ============================================================
# 엑셀 생성
# ============================================================

FONT_NAME = "맑은 고딕"
HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
SAT_FILL = PatternFill("solid", fgColor="F2F2F2")
HOLIDAY_FILL = PatternFill("solid", fgColor="FCE4D6")
LABEL_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build():
    dates = all_dates(YEAR, MONTH)

    wb = Workbook()
    ws = wb.active
    ws.title = "vacation"

    ws["A1"] = f"{YEAR}년 {MONTH}월 휴가/부재 입력표 (vacation)"
    ws["A1"].font = Font(name=FONT_NAME, bold=True, size=13)
    ws["A2"] = "칸에 직접 입력하거나 드롭다운에서 선택 (형식: 구분(종일/오전/오후) + 사유). 빈칸 = 정상 근무"
    ws["A2"].font = Font(name=FONT_NAME, size=9, italic=True)

    header_row1 = 4  # 날짜 숫자
    header_row2 = 5  # 요일
    first_data_row = 6

    ws.cell(row=header_row1, column=1, value="근무자")
    ws.cell(row=header_row2, column=1, value="")
    for r in (header_row1, header_row2):
        c = ws.cell(row=r, column=1)
        c.font = Font(name=FONT_NAME, bold=True, size=10)
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER

    for i, d in enumerate(dates):
        col = 2 + i
        is_sat = d.weekday() == 5
        is_hol = d in HOLIDAYS
        fill = HOLIDAY_FILL if is_hol else (SAT_FILL if is_sat else HEADER_FILL)

        c1 = ws.cell(row=header_row1, column=col, value=d.day)
        c1.font = Font(name=FONT_NAME, bold=True, size=10)
        c1.fill = fill
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c1.border = BORDER

        wd_label = WEEKDAY_NAMES[d.weekday()]
        if is_hol:
            wd_label += f"·{HOLIDAYS[d]}"
        c2 = ws.cell(row=header_row2, column=col, value=wd_label)
        c2.font = Font(name=FONT_NAME, bold=True, size=9)
        c2.fill = fill
        c2.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c2.border = BORDER

        ws.column_dimensions[c1.column_letter].width = 11

    for r, name in enumerate(STAFF_LIST, start=first_data_row):
        c = ws.cell(row=r, column=1, value=name)
        c.font = Font(name=FONT_NAME, bold=True, size=10)
        c.fill = LABEL_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
        for i, d in enumerate(dates):
            col = 2 + i
            cc = ws.cell(row=r, column=col)
            cc.font = Font(name=FONT_NAME, size=10)
            cc.alignment = Alignment(horizontal="center", vertical="center")
            cc.border = BORDER
            if d.weekday() == 5:
                cc.fill = SAT_FILL
            elif d in HOLIDAYS:
                cc.fill = HOLIDAY_FILL

    # 드롭다운(데이터 유효성 검사) 적용
    dv = DataValidation(type="list", formula1='"' + ",".join(CHOICES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    last_row = first_data_row + len(STAFF_LIST) - 1
    last_col_letter = ws.cell(row=first_data_row, column=1 + len(dates)).column_letter
    dv.add(f"B{first_data_row}:{last_col_letter}{last_row}")

    ws.column_dimensions["A"].width = 10
    ws.freeze_panes = ws.cell(row=first_data_row, column=2)

    base_name = "vacation"
    path = os.path.join(OUTPUT_DIR, base_name + ".xlsx")
    n = 1
    while os.path.exists(path):
        path = os.path.join(OUTPUT_DIR, f"{base_name}-{n}.xlsx")
        n += 1
    wb.save(path)
    return path


if __name__ == "__main__":
    path = build()
    print(f"저장 완료: {path}")
